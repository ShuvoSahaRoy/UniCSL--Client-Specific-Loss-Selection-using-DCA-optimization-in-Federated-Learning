import sys, os, random, numpy as np, pandas as pd, cv2, torch, torch.nn as nn, torch.nn.functional as F, itertools
from datetime import datetime
import matplotlib.pyplot as plt
import seaborn as sns, pickle
from openpyxl import load_workbook, Workbook
from openpyxl.utils import get_column_letter

from config import *
# num_clients = 100
# non_iid = 1
# save = 0
# CR = 30
# SEED = 1234
# corrupt_data = 0
# val_fraction = 0.1
# results_file =  f'results/{SEED}_experimental_results.xlsx'

class LinearSVM_OVO(nn.Module):
    def __init__(self, input_dim: int, num_classes: int):
        super().__init__()
        self.num_classes = num_classes

        # All (i, j) pairs with i < j
        pairs = list(itertools.combinations(range(num_classes), 2))
        self.num_pairs = len(pairs)

        self.linear = nn.Linear(input_dim, self.num_pairs)

        # --- OvO pair tensors (REGISTERED BUFFERS) ---
        pair_i = torch.tensor([i for i, j in pairs], dtype=torch.long)
        pair_j = torch.tensor([j for i, j in pairs], dtype=torch.long)

        self.register_buffer("pair_i", pair_i)
        self.register_buffer("pair_j", pair_j)

    def forward(self, x):
        return self.linear(x)  

def set_seed(SEED):
    random.seed(SEED)
    np.random.seed(SEED)
    torch.manual_seed(SEED)

def generate_client_schedule(num_clients, participants, rounds):
    return [np.random.choice(num_clients, size=participants, replace=False) for _ in range(rounds)]

class Tee:
    def __init__(self, *streams):
        self.streams = streams
        self.original_stdout = sys.stdout
        self.original_stderr = sys.stderr

    def write(self, message):
        for s in self.streams:
            s.write(message)
            s.flush()

    def flush(self):
        for s in self.streams:
            s.flush()

    def close(self):
        sys.stdout = self.original_stdout
        sys.stderr = self.original_stderr

def setup_logger(save_log=False, log_dir="results/logs"):
    if save_log:
        os.makedirs(log_dir, exist_ok=True)
        timestamp = datetime.now().strftime("%m-%d_%I-%M-%S-%p")
        log_file_path = os.path.join(log_dir, f"experiment_log_{SEED}_{timestamp}.txt")
        log_file = open(log_file_path, "w")
        tee = Tee(sys.stdout, log_file)
        sys.stdout = tee
        sys.stderr = tee
        return log_file, tee
    else:
        return None, None



# Define the ordered loss function names
loss_functions_ordered = [
    'Least Squares', 'Smooth Hinge', 'Squared Hinge', 'Truncated SH', 'Truncated LS',
    'Smoothed Ramp1', 'Smoothed Ramp2', 'nonconvex exp loss1', 'nonconvex exp loss2', 'nonconvex exp loss3',
]

def plot_and_write_function_frequencies(loss_dictionary, dataset_name):
    # Step 1: Initialize loss count with zero for all loss functions
    loss_count = {loss_name: 0 for loss_name in loss_functions_ordered}

    # Step 2: Count the actual frequency from the loss_dictionary
    for key, value in loss_dictionary.items():
        if value.get('best_loss') is None:
            continue
        for loss_function in value['best_loss'].keys():
            if loss_function in loss_count:
                loss_count[loss_function] += 1

    df_new = pd.DataFrame({
        'Dataset Name': [dataset_name] * len(loss_functions_ordered),
        'Loss Function': loss_functions_ordered,
        'Frequency': [loss_count[loss] for loss in loss_functions_ordered]
    })

    # Step 5: Append the data to the main experimental results Excel file
    if save:
        # choose sheet name based on iid/noniid
        if non_iid == 0:
            sheet_name = 'loss_freq_iid'
        elif non_iid == 1 and not corrupt_data:
            sheet_name = 'loss_freq_non_iid'
        elif non_iid == 1 and corrupt_data:
            sheet_name = 'loss_freq_non_iid_corrupted'
        else:
            raise ValueError("Unsupported configuration")

        os.makedirs(os.path.dirname(results_file) if os.path.dirname(results_file) else '.', exist_ok=True)

        if os.path.exists(results_file):
            wb = load_workbook(results_file)
            if sheet_name in wb.sheetnames:
                sheet = wb[sheet_name]
            else:
                sheet = wb.create_sheet(sheet_name)
                # write header
                sheet.append(list(df_new.columns))

            # If header missing (empty sheet), ensure header exists
            if sheet.max_row == 1 and sheet['A1'].value is None:
                sheet.append(list(df_new.columns))

            # Append rows
            for _, row in df_new.iterrows():
                sheet.append([row['Dataset Name'], row['Loss Function'], int(row['Frequency']) if not pd.isna(row['Frequency']) else row['Frequency']])

            wb.save(results_file)
            wb.close()
        else:
            wb = Workbook()
            sheet = wb.active
            sheet.title = sheet_name
            # write header
            sheet.append(list(df_new.columns))
            # append rows
            for _, row in df_new.iterrows():
                sheet.append([row['Dataset Name'], row['Loss Function'], int(row['Frequency']) if not pd.isna(row['Frequency']) else row['Frequency']])
            wb.save(results_file)
            wb.close()

    return df_new


# ---------------------------------------
# Helper: map (noise, non_iid) ➜ suffix
# ---------------------------------------
def scenario_suffix(noise: int, non_iid: int) -> str:
    if noise == 0 and non_iid == 0:
        return "iid"
    if noise == 0 and non_iid == 1:
        return "noniid"
    if noise == 1 and non_iid == 0:
        return "noiseiid"
    return "noisenoniid"          # noise ==1 and non_iid ==1


def plot_heatmap(df, dataset, suffix, max_value, font_sizes, xtick_interval, ytick_interval):
    plt.figure(figsize=(12, 8))
    ax = sns.heatmap(
        df, cmap='YlOrRd', annot=False,
        vmin=0, vmax=max_value,
        xticklabels=xtick_interval, yticklabels=ytick_interval,
        cbar_kws={'label': 'Local Epochs'}
    )
    cbar = ax.collections[0].colorbar
    ax.tick_params(axis='x', labelsize=18)
    ax.tick_params(axis='y', labelsize=18)
    cbar.set_label('Local Epochs', fontsize=font_sizes['cbar_label'])
    cbar.ax.tick_params(labelsize=font_sizes['cbar_ticks'])
    
    plt.xlabel('Communication Rounds', fontsize=font_sizes['xlabel'])
    plt.ylabel('Clients', fontsize=font_sizes['ylabel'])
    plt.tight_layout()
    if save:
        plt.savefig(f'results/{dataset}_{suffix}_heatmap.png', dpi=300)
    plt.close()

def heatmap_gen(multi_client_iter_data, single_client_iter_data, dataset):
    # Default font sizes
    font_sizes = {
        'cbar_label': 24,
        'cbar_ticks': 22,
        'xlabel': 24,
        'ylabel': 24
    }

    xtick_interval = 5
    ytick_interval = 10
    # Prepare DataFrames
    df_multi = pd.DataFrame(
        multi_client_iter_data,
        index=[f'Client_{i}' for i in range(num_clients)],
        columns=[f'{i}' for i in range(CR)]
    )
    df_single = pd.DataFrame(
        single_client_iter_data,
        index=[f'Client_{i}' for i in range(num_clients)],
        columns=[f'{i}' for i in range(CR)]
    )

    # Determine heatmap color scale max
    max_value = max(df_multi.max().max(), df_single.max().max())

    # Plot heatmaps
    plot_heatmap(df_multi, dataset, 'multi_loss', max_value, font_sizes, xtick_interval, ytick_interval)
    plot_heatmap(df_single, dataset, 'single_loss', max_value, font_sizes, xtick_interval, ytick_interval)

def save_accuracy_metrics_to_excel(algo_name, metric_tuple, dataset_name, filename=results_file):

    if not save:
        return
    accuracy_metric = metric_tuple[0]
    total_time = metric_tuple[1]
    peak_memory = metric_tuple[2]

    # --------------------------------------------------
    # Determine sheet base
    # --------------------------------------------------
    if non_iid == 0:
        base_sheet = "iid"
    elif non_iid == 1 and not corrupt_data:
        base_sheet = "non_iid"
    elif non_iid == 1 and corrupt_data==1:
        base_sheet = "non_iid_corrupted"
    else:
        raise ValueError("Unsupported configuration")

    acc_sheet = base_sheet
    time_sheet = base_sheet + "_time"
    memory_sheet = base_sheet + "_memory"

    # --------------------------------------------------
    # Prepare accuracy list
    # --------------------------------------------------
    if isinstance(accuracy_metric, (list, np.ndarray)):
        accuracy_list = [round(float(acc), 2) for acc in accuracy_metric]
    else:
        accuracy_list = [round(float(accuracy_metric), 2)]

    rounds = list(range(1, len(accuracy_list) + 1))

    # --------------------------------------------------
    # Build config string
    # --------------------------------------------------
    config = (
        f"clients={num_clients}, non_iid={non_iid}, "
        f"alpha={alpha}, seed={SEED}, corrupt={corrupt_data}, mu = {mu}, alpha_coef = {alpha_coef}"
    )

    os.makedirs(os.path.dirname(filename), exist_ok=True)

    # --------------------------------------------------
    # 1️⃣ SAVE ACCURACY (per round + average)
    # --------------------------------------------------
    df_acc_new = pd.DataFrame({
        "Config": [config] * len(rounds),
        "Dataset": [dataset_name] * len(rounds),
        "Round": rounds,
        algo_name: accuracy_list
    })
    
    # Add average row
    avg_accuracy = round(float(np.mean(accuracy_list)), 2)
    df_avg = pd.DataFrame({
        "Config": [config],
        "Dataset": [dataset_name],
        "Round": ["Average"],
        algo_name: [avg_accuracy]
    })
    
    df_acc_new = pd.concat([df_acc_new, df_avg], ignore_index=True)

    _merge_and_save(df_acc_new, acc_sheet, ["Config", "Dataset", "Round"], algo_name, filename)

    # ==================================================
    # 2️⃣ SAVE EXECUTION TIME
    # ==================================================
    df_time_new = pd.DataFrame({
        "Config": [config],
        "Dataset": [dataset_name],
        algo_name: [round(float(total_time), 4)]
    })

    _merge_and_save(df_time_new, time_sheet, ["Config", "Dataset"], algo_name, filename)

    # ==================================================
    # 3️⃣ SAVE PEAK MEMORY
    # ==================================================
    df_mem_new = pd.DataFrame({
        "Config": [config],
        "Dataset": [dataset_name],
        algo_name: [round(float(peak_memory), 4)]
    })

    _merge_and_save(df_mem_new, memory_sheet, ["Config", "Dataset"], algo_name, filename)

    print(f"Saved {algo_name} results successfully.")

def _merge_and_save(df_new, sheet_name, index_cols, algo_name, filename):

    # Load existing sheet
    if os.path.exists(filename):
        try:
            df_old = pd.read_excel(filename, sheet_name=sheet_name)
        except ValueError:
            df_old = pd.DataFrame()
    else:
        df_old = pd.DataFrame()

    if df_old.empty:
        final_df = df_new
    else:
        final_df = pd.merge(
            df_old,
            df_new,
            on=index_cols,
            how="outer"
        )

        # If same algorithm already exists, prefer new values
        if f"{algo_name}_x" in final_df.columns:
            final_df[algo_name] = final_df[f"{algo_name}_y"].combine_first(
                final_df[f"{algo_name}_x"]
            )
            final_df.drop(columns=[f"{algo_name}_x", f"{algo_name}_y"], inplace=True)

    # Reorder index columns first
    cols = final_df.columns.tolist()
    for col in reversed(index_cols):
        cols.remove(col)
        cols.insert(0, col)

    final_df = final_df[cols]

    # Save safely
    mode = "a" if os.path.exists(filename) else "w"
    if_sheet = "replace" if mode == "a" else None

    with pd.ExcelWriter(filename,
                        engine="openpyxl",
                        mode=mode,
                        if_sheet_exists=if_sheet) as writer:
        final_df.to_excel(writer, sheet_name=sheet_name, index=False)


"""Part of the code below is for data corruption and visualization of corrupted samples, 
which is used for the noise experiments. It is not part of the main training loop but can
be used to generate and visualize corrupted data samples."""
def add_gaussian_noise(img, std=0.3):
    noise = np.random.normal(0, std, img.shape)
    return np.clip(img + noise, -1, 1)

def add_salt_pepper_noise(img, amount=0.2):
    noisy = img.copy()
    num_pixels = int(amount * img.size)
    coords = np.random.randint(0, img.shape[0], (num_pixels, 2))
    for y, x in coords:
        noisy[y, x] = np.random.choice([-1, 1])
    return noisy

def add_blur(img):
    return cv2.GaussianBlur(img, (5, 5), 0)

def add_occlusion(img):
    img = img.copy()
    h, w = img.shape
    img[h//4:h//2, w//4:w//2] = 0
    return img

def add_label_flip(label):
    new_label = np.random.randint(0, 10)
    while new_label == label:
        new_label = np.random.randint(0, 10)
    return new_label

def corrupt_client_data(train_df, noise_type, corruption_rate=0.2):

    n_samples = len(train_df)
    n_corrupt = int(corruption_rate * n_samples)

    corrupt_indices = np.random.choice(n_samples, n_corrupt, replace=False)

    for idx in corrupt_indices:
        row = train_df.iloc[idx]

        if noise_type == "label_flip":
            train_df.at[idx, "y"] = add_label_flip(row["y"])
            continue

        image = row.drop("y").values.reshape(28, 28)

        if noise_type == "gaussian":
            image = add_gaussian_noise(image)

        elif noise_type == "salt_pepper":
            image = add_salt_pepper_noise(image)

        elif noise_type == "blur":
            image = add_blur(image)

        elif noise_type == "occlusion":
            image = add_occlusion(image)

        train_df.loc[idx, train_df.columns != "y"] = image.flatten().astype(np.float32)

    return train_df




def visualize_multiple_samples(train_data_list, client_id, n_samples=5):
    
    client_df = train_data_list[client_id]
    indices = np.random.choice(len(client_df), n_samples, replace=False)

    plt.figure(figsize=(10, 2))
    
    for i, idx in enumerate(indices):
        row = client_df.iloc[idx]
        image = row.drop("y").values.reshape(28, 28)
        label = int(row["y"])

        plt.subplot(1, n_samples, i+1)
        plt.imshow(image, cmap="gray")
        plt.title(label)
        plt.axis("off")

    plt.suptitle(f"Client {client_id}")
    plt.tight_layout()
    plt.show()


def visualize_corruption_examples(train_df):

    corruption_types = ["gaussian", "salt_pepper", "blur", "occlusion", "label_flip"]

    # Pick one fixed sample
    sample = train_df.iloc[0].copy()
    original_image = sample.drop("y").values.reshape(28, 28)
    original_label = int(sample["y"])

    plt.figure(figsize=(14, 3))

    # ---- Column 1: Original ----
    plt.subplot(1, 6, 1)
    plt.imshow(original_image, cmap="gray")
    plt.title(f"Original\nLabel:{original_label}")
    plt.axis("off")

    # ---- Columns 2–6: Corruptions ----
    for i, noise_type in enumerate(corruption_types):

        temp_df = sample.to_frame().T.copy()

        corrupted_df = corrupt_client_data(temp_df,noise_type,corruption_rate=1.0)

        corrupted_sample = corrupted_df.iloc[0]
        corrupted_image = corrupted_sample.drop("y").values.reshape(28, 28)
        corrupted_label = int(corrupted_sample["y"])

        plt.subplot(1, 6, i+2)

        if noise_type == "label_flip":
            title = f"Label Flip\n{original_label} → {corrupted_label}"
        else:
            title = noise_type.capitalize()

        plt.imshow(corrupted_image, cmap="gray")
        plt.title(title)
        plt.axis("off")

    plt.tight_layout()
    plt.show()



def load_lr_rate(dataset, algorithm, non_iid, filename="lr_rate.xlsx"):
    """
    Load tuned learning rate based on dataset, algorithm, and scenario.
    
    Args:
        dataset (str): dataset name (e.g., 'mnist')
        algorithm (str): algorithm name (e.g., 'FedAvg')
        non_iid (int): 0 -> iid sheet, 1 -> noniid sheet
        filename (str): excel file path
    
    Returns:
        float: learning rate
    """

    # Select sheet based on scenario
    sheet_name = "noniid" if non_iid == 1 else "iid"

    # Read sheet
    df = pd.read_excel(filename, sheet_name=sheet_name)

    # Normalize dataset column
    df["Dataset"] = df["Dataset"].str.lower()
    dataset = dataset.lower()

    # Check dataset exists
    if dataset not in df["Dataset"].values:
        raise ValueError(f"Dataset '{dataset}' not found in {sheet_name} sheet.")

    # Check algorithm exists
    if algorithm not in df.columns:
        raise ValueError(f"Algorithm '{algorithm}' not found in {sheet_name} sheet.")

    # Extract LR
    lr = df.loc[df["Dataset"] == dataset, algorithm].values[0]

    return float(lr)

# for unicsl
def load_best_c(excel_path, seed, dataset, non_iid, corrupt_data, c_list=[0.01, 0.1, 1, 10, 20, 30]):
    """
    Returns the best C value (highest mean accuracy) for given settings.

    Parameters
    ----------
    excel_path : str
        Path to the Excel results file
    seed : int
    dataset : str
    non_iid : int
    corrupt_data : int
    c_list : list of float
        List of C values used in tuning

    Returns
    -------
    best_c : float
    best_acc : float
    """

    df = pd.read_excel(excel_path)

    scenario = f"{non_iid},{corrupt_data}"

    # filter row
    row = df[
        (df["seed"] == seed) &
        (df["dataset"] == dataset) &
        (df["scenario"] == scenario)
    ]

    if row.empty:
        raise ValueError("No matching row found for given parameters.")

    row = row.iloc[0]

    best_c = None
    best_acc = -np.inf

    for c in c_list:
        col = f"c = {c}"

        if col not in row:
            continue

        cell_value = row[col]

        if isinstance(cell_value, str) and "±" in cell_value:
            mean_str = cell_value.split("±")[0].strip()
            try:
                mean_val = float(mean_str)
            except:
                continue

            if mean_val > best_acc:
                best_acc = mean_val
                best_c = c

    if best_c is None:
        raise ValueError("Could not determine best C.")

    return best_c, best_acc