import re
from openpyxl import load_workbook

# ----------------------------
# Load Excel
# ----------------------------
file = "confidence_intervals.xlsx"      # change path if needed
wb = load_workbook(file, data_only=True)
ws = wb.active

# ----------------------------
# Result dictionary
# ----------------------------
summary = {
    "Tabular": {"UniCSL": 0, "Baseline": 0, "NS": 0},
    "Image":   {"UniCSL": 0, "Baseline": 0, "NS": 0},
}

# ----------------------------
# Current section
# ----------------------------
current_type = None

# Regular expression for extracting CI
pattern = re.compile(r"\(([-+]?\d*\.?\d+),\s*([-+]?\d*\.?\d+)\)")

for row in ws.iter_rows(values_only=True):

    first = row[0]

    if first is None:
        continue

    # Detect dataset category
    if first == "Tabular":
        current_type = "Tabular"
        continue

    if first == "Image":
        current_type = "Image"
        continue

    # Skip titles/header rows
    if first in [
        "Dataset",
        "IID",
        "NonIID",
        "Non-IID with Noise",
        "Mean Accuracy Difference and 95% Confidence Interval (Baseline − UniCSL)"
    ]:
        continue

    # Dataset row
    if current_type is None:
        continue

    # columns B:G
    for cell in row[1:7]:

        if cell is None:
            continue

        m = pattern.search(str(cell))
        if not m:
            continue

        lower = float(m.group(1))
        upper = float(m.group(2))

        if upper < 0:
            summary[current_type]["UniCSL"] += 1

        elif lower > 0:
            summary[current_type]["Baseline"] += 1

        else:
            summary[current_type]["NS"] += 1


# ----------------------------
# Print summary
# ----------------------------
overall = {"UniCSL":0,"Baseline":0,"NS":0}

print("="*70)
print(f"{'Dataset Type':<15}{'UniCSL':>10}{'Baseline':>12}{'No Sig.':>12}{'Total':>10}")
print("="*70)

for t in ["Tabular","Image"]:

    total = (
        summary[t]["UniCSL"]
        + summary[t]["Baseline"]
        + summary[t]["NS"]
    )

    overall["UniCSL"] += summary[t]["UniCSL"]
    overall["Baseline"] += summary[t]["Baseline"]
    overall["NS"] += summary[t]["NS"]

    print(
        f"{t:<15}"
        f"{summary[t]['UniCSL']:>10}"
        f"{summary[t]['Baseline']:>12}"
        f"{summary[t]['NS']:>12}"
        f"{total:>10}"
    )

overall_total = overall["UniCSL"] + overall["Baseline"] + overall["NS"]

print("="*70)
print(
    f"{'Overall':<15}"
    f"{overall['UniCSL']:>10}"
    f"{overall['Baseline']:>12}"
    f"{overall['NS']:>12}"
    f"{overall_total:>10}"
)
print("="*70)