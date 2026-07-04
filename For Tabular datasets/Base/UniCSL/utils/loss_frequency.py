import pandas as pd
import os
import matplotlib.pyplot as plt
from openpyxl import load_workbook

# def plot_and_write_function_frequencies(loss_dictionary, dataset_name, plot=False, store=False, excel_file='loss_frequencies.xlsx'):
#     # Step 1: Count the frequency of each loss function
#     loss_count = {}
#     for key, value in loss_dictionary.items():
#         for loss_function in value['best_loss'].keys():
#             if loss_function in loss_count:
#                 loss_count[loss_function] += 1
#             else:
#                 loss_count[loss_function] = 1

#     # Step 2: Plot the frequency graph
#     if plot:
#         plt.figure(figsize=(10, 6))
#         plt.bar(loss_count.keys(), loss_count.values(), color='skyblue')
#         plt.xlabel('Loss Function')
#         plt.ylabel('Frequency')
#         plt.title('Frequency of Selected Loss Functions')
#         plt.xticks(rotation=45, ha='right')
#         plt.tight_layout()
#         plt.show()

#     # Step 3: Create a new DataFrame with dataset_name, loss_function, and frequency
#     df_new = pd.DataFrame({
#         'Dataset Name': [dataset_name] * len(loss_count),
#         'Loss Function': list(loss_count.keys()),
#         'Frequency': list(loss_count.values())
#     })

#     # Step 4: Append the data to an existing Excel file or create a new one
#     if store:
#         if os.path.exists(excel_file):
#             # Load the existing workbook
#             with pd.ExcelWriter(excel_file, engine='openpyxl', mode='a', if_sheet_exists='overlay') as writer:
#                 # Get the current data and write without header
#                 df_new.to_excel(writer, index=False, header=False, startrow=writer.sheets['Sheet1'].max_row)
#         else:
#             # If the file doesn't exist, create a new one and write the data with headers
#             df_new.to_excel(excel_file, index=False)
    
#     return df_new



# Define the ordered loss function names
loss_functions_ordered = [
    'Least Squares', 'Smooth Hinge', 'Squared Hinge', 'Truncated SH', 'Truncated LS',
    'Smoothed Ramp1', 'Smoothed Ramp2', 'nonconvex exp loss1', 'nonconvex exp loss2', 'nonconvex exp loss3'
]

def plot_and_write_function_frequencies(loss_dictionary, dataset_name, plot=False, store=False, excel_file='loss_frequencies.xlsx'):
    # Step 1: Initialize loss count with zero for all loss functions
    loss_count = {loss_name: 0 for loss_name in loss_functions_ordered}

    # Step 2: Count the actual frequency from the loss_dictionary
    for key, value in loss_dictionary.items():
        for loss_function in value['best_loss'].keys():
            if loss_function in loss_count:
                loss_count[loss_function] += 1

    # Step 3: Plot the frequency graph
    if plot:
        plt.figure(figsize=(10, 6))
        plt.bar(loss_count.keys(), loss_count.values(), color='skyblue')
        plt.xlabel('Loss Function')
        plt.ylabel('Frequency')
        plt.title(f'Frequency of Selected Loss Functions ({dataset_name})')
        plt.xticks(rotation=45, ha='right')
        plt.tight_layout()
        plt.show()

    # Step 4: Create a new DataFrame following the fixed order
    df_new = pd.DataFrame({
        'Dataset Name': [dataset_name] * len(loss_functions_ordered),
        'Loss Function': loss_functions_ordered,
        'Frequency': [loss_count[loss] for loss in loss_functions_ordered]
    })

    # Step 5: Append the data to an existing Excel file or create a new one
    if store:
        if os.path.exists(excel_file):
            # Load the existing workbook and append the data without header
            with pd.ExcelWriter(excel_file, engine='openpyxl', mode='a', if_sheet_exists='overlay') as writer:
                df_new.to_excel(writer, index=False, header=False, startrow=writer.sheets['Sheet1'].max_row)
        else:
            # Create a new file and write data with headers
            df_new.to_excel(excel_file, index=False)

    return df_new