from Base.utils import corrupt_data
import sys, os, random, torch, pandas as pd, numpy as np
from datetime import datetime
import matplotlib.pyplot as plt
import seaborn as sns
from openpyxl import load_workbook
import time, psutil

from config import num_clients, min_samples_per_client, alpha, non_iid, CR, noise, participants, save, SEED, local_epoch

def SERVER(global_model, local_updates, n):
 
    global_model = np.sum([w * (n_i / n) for w, n_i in local_updates], axis=0)
    
    # delta = np.sum([w * (n_i / n) for w, n_i in local_updates], axis=0)
    # global_model = global_model - lr * delta
    # global_model = global_model/np.linalg.norm(global_model) 

    return global_model


def set_seed():
    random.seed(SEED)
    np.random.seed(SEED)
    torch.manual_seed(SEED)


class Tee:
    def __init__(self, *streams):
        self.streams = streams
        self.original_stdout = sys.stdout
        self.original_stderr = sys.stderr

    def write(self, message):
        for s in self.streams:
            s.write(message)
            s.flush()

    def flush(self):
        for s in self.streams:
            s.flush()

    def close(self):
        sys.stdout = self.original_stdout
        sys.stderr = self.original_stderr

def setup_logger(save_log=False, log_dir="results/logs"):
    if save_log:
        os.makedirs(log_dir, exist_ok=True)
        timestamp = datetime.now().strftime("%Y-%m-%d_%H-%M-%S")
        log_file_path = os.path.join(log_dir, f"experiment_log_{timestamp}.txt")
        log_file = open(log_file_path, "w")
        tee = Tee(sys.stdout, log_file)
        sys.stdout = tee
        sys.stderr = tee
        return log_file, tee
    else:
        return None, None

def load_lr_rate(dataset, algorithm, non_iid, filename="lr_rate.xlsx"):
    algorithm = algorithm.split('_')[0]
    """
    Load tuned learning rate based on dataset, algorithm, and scenario.
    
    Args:
        dataset (str): dataset name (e.g., 'mnist')
        algorithm (str): algorithm name (e.g., 'FedAvg')
        non_iid (int): 0 -> iid sheet, 1 -> noniid sheet
        filename (str): excel file path
    
    Returns:
        float: learning rate
    """

    # Select sheet based on scenario
    sheet_name = "noniid_sq" if non_iid == 1 else "iid_sq"

    # Read sheet
    df = pd.read_excel(filename, sheet_name=sheet_name)

    # Normalize dataset column
    df["dataset"] = df["dataset"].str.lower()
    dataset = dataset.lower()

    # Check dataset exists
    if dataset not in df["dataset"].values:
        raise ValueError(f"Dataset '{dataset}' not found in {sheet_name} sheet.")

    # Check algorithm exists
    if algorithm not in df.columns:
        raise ValueError(f"Algorithm '{algorithm}' not found in {sheet_name} sheet.")

    # Extract LR
    lr = df.loc[df["dataset"] == dataset, algorithm].values[0]

    return float(lr)

# try this noniid for 5.4 experiments for same results.
def noniid_distribution(data):
    total_samples = len(data)
    labels = data.iloc[:, -1]
    class_0_indices = data[labels == 0].index.to_numpy().copy()
    class_1_indices = data[labels == 1].index.to_numpy().copy()
    np.random.shuffle(class_0_indices)
    np.random.shuffle(class_1_indices)

    min_samples_assigned = min_samples_per_client * num_clients
    if min_samples_assigned > total_samples:
        raise ValueError("Not enough samples for minimum allocation.")
    
    samples_per_class = min_samples_per_client // 2

    if len(class_0_indices) < num_clients * samples_per_class or len(class_1_indices) < num_clients * samples_per_class:
        raise ValueError("Not enough samples in one or both classes for minimum allocation.")

    client_splits = [
        np.concatenate([
            class_0_indices[i * samples_per_class:(i + 1) * samples_per_class],
            class_1_indices[i * samples_per_class:(i + 1) * samples_per_class]
        ]) for i in range(num_clients)
    ]

    remaining_indices = np.concatenate([
        class_0_indices[num_clients * samples_per_class:],
        class_1_indices[num_clients * samples_per_class:]
    ])
    np.random.shuffle(remaining_indices)
    remaining_samples = len(remaining_indices)

    if remaining_samples > 0:
        max_samples_per_client = remaining_samples // 2 
        proportions = np.random.dirichlet([alpha] * num_clients)
        split_sizes = np.minimum((proportions * remaining_samples).astype(int), max_samples_per_client)
        
        deficit = remaining_samples - split_sizes.sum()
        if deficit > 0:
            sorted_indices = np.argsort(split_sizes)
            for i in sorted_indices:
                additional = min(max_samples_per_client - split_sizes[i], deficit)
                split_sizes[i] += additional
                deficit -= additional
                if deficit == 0:
                    break
        
        split_points = np.cumsum(split_sizes)[:-1]
        split_points = split_points[split_points < remaining_samples]  
        remaining_client_splits = np.split(remaining_indices, split_points) if split_points.size > 0 else [remaining_indices]
        
        for i in range(num_clients):
            if i < len(remaining_client_splits):
                client_splits[i] = np.concatenate([client_splits[i], remaining_client_splits[i]])

    # Create dataframes and verify total samples (issue 6)
    client_dataframes = [data.loc[split].reset_index(drop=True) for split in client_splits]
    assert sum(len(df) for df in client_dataframes) == len(data), "Sample mismatch"
    return client_dataframes

# corrected implementation
# def noniid_distribution(data, num_clients=num_clients, alpha=alpha, min_samples_per_client=min_samples_per_client):

#     total_samples = len(data)

#     if min_samples_per_client * num_clients > total_samples:
#         raise ValueError("Not enough samples for minimum allocation.")

#     # ── Step 1: Random minimum allocation ───────────────────────────────
#     all_indices = data.index.to_numpy().copy()
#     np.random.shuffle(all_indices)

#     client_splits = []
#     offset = 0

#     for _ in range(num_clients):
#         client_splits.append(
#             all_indices[offset:offset + min_samples_per_client]
#         )
#         offset += min_samples_per_client

#     # Remaining samples
#     remaining_indices = all_indices[offset:]
#     remaining_data = data.loc[remaining_indices]

#     labels = remaining_data.iloc[:, -1]

#     class_0_indices = remaining_data[labels == 0].index.to_numpy().copy()
#     class_1_indices = remaining_data[labels == 1].index.to_numpy().copy()

#     np.random.shuffle(class_0_indices)
#     np.random.shuffle(class_1_indices)

#     # ── Step 2: per-class Dirichlet split of the REMAINING pool ──────────────

#     remaining_class_0 = class_0_indices
#     remaining_class_1 = class_1_indices

#     def _dirichlet_split(class_indices, num_clients, alpha):
#         """Split one class's indices across clients via Dirichlet proportions."""
#         n_remaining = len(class_indices)
#         if n_remaining == 0:
#             return [np.array([], dtype=class_indices.dtype) for _ in range(num_clients)]

#         proportions = np.random.dirichlet([alpha] * num_clients)
#         split_sizes = (proportions * n_remaining).astype(int)

#         # Fix rounding deficit so all samples are assigned (none dropped)
#         deficit = n_remaining - split_sizes.sum()
#         if deficit > 0:
#             # distribute leftover one-by-one to clients with largest
#             # fractional remainder (more stable than always favoring the same client)
#             fractional = proportions * n_remaining - split_sizes
#             order = np.argsort(-fractional)   # descending fractional part
#             for i in order[:deficit]:
#                 split_sizes[i] += 1

#         split_points = np.cumsum(split_sizes)[:-1]
#         return np.split(class_indices, split_points)

#     class_0_splits = _dirichlet_split(remaining_class_0, num_clients, alpha)
#     class_1_splits = _dirichlet_split(remaining_class_1, num_clients, alpha)

#     # ── Merge minimum floor + Dirichlet-skewed remainder per client ──────────
#     for i in range(num_clients):
#         extra = np.concatenate([class_0_splits[i], class_1_splits[i]])
#         if extra.size > 0:
#             np.random.shuffle(extra)
#             client_splits[i] = np.concatenate([client_splits[i], extra])

#     client_dataframes = [data.loc[split].reset_index(drop=True) for split in client_splits]

#     assert sum(len(df) for df in client_dataframes) == len(data), "Sample mismatch"
#     return client_dataframes



def iid_distribution(data):
    # Get the label column (last column)
    labels = data.iloc[:, -1]
    # Group data by labels
    grouped = data.groupby(labels)
    
    # Initialize empty lists for each client's data
    client_datasets = [[] for _ in range(num_clients)]
    
    # Split each class's data into num_clients parts
    for _, group in grouped:
        # Split the group's indices into num_clients parts
        split_indices = np.array_split(group.index, num_clients)
        # Append each split to the corresponding client's dataset
        for i, indices in enumerate(split_indices):
            client_datasets[i].append(group.loc[indices])
    
    # Concatenate each client's splits into a single DataFrame
    client_datasets = [pd.concat(splits, ignore_index=True) for splits in client_datasets]
    
    # Verify no samples are lost
    total_samples = sum(len(client_data) for client_data in client_datasets)
    assert total_samples == len(data), f"Sample mismatch: {total_samples} != {len(data)}"
    
    return client_datasets



def save_experimental_results(dataset_name, algo_name, metric_tuple, save, filename=f'results/{SEED}_experimental_results.xlsx'):

    if not save:
        return
    accuracy_metric = metric_tuple[0]
    total_time = metric_tuple[1]
    peak_memory = metric_tuple[2]

    # --------------------------------------------------
    # Determine sheet base
    # --------------------------------------------------
    if non_iid == 0:
        base_sheet = "iid"
    elif non_iid == 1 and not noise:
        base_sheet = "non_iid"
    elif non_iid == 1 and noise==1:
        base_sheet = "non_iid_corrupted"
    else:
        raise ValueError("Unsupported configuration")

    acc_sheet = base_sheet
    time_sheet = base_sheet + "_time"
    memory_sheet = base_sheet + "_memory"

    # --------------------------------------------------
    # Prepare accuracy list
    # --------------------------------------------------
    if isinstance(accuracy_metric, (list, np.ndarray)):
        accuracy_list = [round(float(acc), 2) for acc in accuracy_metric]
    else:
        accuracy_list = [round(float(accuracy_metric), 2)]

    rounds = list(range(1, len(accuracy_list) + 1))

    # --------------------------------------------------
    # Build config string
    # --------------------------------------------------
    config = (
        f"clients={num_clients}, non_iid={non_iid}, "
        f"alpha={alpha}, seed={SEED}, corrupt={noise}"
    )

    os.makedirs(os.path.dirname(filename), exist_ok=True)

    # --------------------------------------------------
    # 1️⃣ SAVE ACCURACY (per round + average)
    # --------------------------------------------------
    df_acc_new = pd.DataFrame({
        "Config": [config] * len(rounds),
        "Dataset": [dataset_name] * len(rounds),
        "Round": rounds,
        algo_name: accuracy_list
    })
    
    # Add average row
    avg_accuracy = round(float(np.mean(accuracy_list)), 2)
    df_avg = pd.DataFrame({
        "Config": [config],
        "Dataset": [dataset_name],
        "Round": ["Average"],
        algo_name: [avg_accuracy]
    })
    
    df_acc_new = pd.concat([df_acc_new, df_avg], ignore_index=True)

    _merge_and_save(df_acc_new, acc_sheet, ["Config", "Dataset", "Round"], algo_name, filename)

    # ==================================================
    # 2️⃣ SAVE EXECUTION TIME
    # ==================================================
    df_time_new = pd.DataFrame({
        "Config": [config],
        "Dataset": [dataset_name],
        algo_name: [round(float(total_time), 4)]
    })

    _merge_and_save(df_time_new, time_sheet, ["Config", "Dataset"], algo_name, filename)

    # ==================================================
    # 3️⃣ SAVE PEAK MEMORY
    # ==================================================
    df_mem_new = pd.DataFrame({
        "Config": [config],
        "Dataset": [dataset_name],
        algo_name: [round(float(peak_memory), 4)]
    })

    _merge_and_save(df_mem_new, memory_sheet, ["Config", "Dataset"], algo_name, filename)

    print(f"Saved {algo_name} results successfully.")

def _merge_and_save(df_new, sheet_name, index_cols, algo_name, filename):

    # Load existing sheet
    if os.path.exists(filename):
        try:
            df_old = pd.read_excel(filename, sheet_name=sheet_name)
        except ValueError:
            df_old = pd.DataFrame()
    else:
        df_old = pd.DataFrame()

    if df_old.empty:
        final_df = df_new
    else:
        final_df = pd.merge(
            df_old,
            df_new,
            on=index_cols,
            how="outer"
        )

        # If same algorithm already exists, prefer new values
        if f"{algo_name}_x" in final_df.columns:
            final_df[algo_name] = final_df[f"{algo_name}_y"].combine_first(
                final_df[f"{algo_name}_x"]
            )
            final_df.drop(columns=[f"{algo_name}_x", f"{algo_name}_y"], inplace=True)

    # Reorder index columns first
    cols = final_df.columns.tolist()
    for col in reversed(index_cols):
        cols.remove(col)
        cols.insert(0, col)

    final_df = final_df[cols]

    # Save safely
    mode = "a" if os.path.exists(filename) else "w"
    if_sheet = "replace" if mode == "a" else None

    with pd.ExcelWriter(filename,
                        engine="openpyxl",
                        mode=mode,
                        if_sheet_exists=if_sheet) as writer:
        final_df.to_excel(writer, sheet_name=sheet_name, index=False)


# Define the ordered loss function names
loss_functions_ordered = [
    'Least Squares', 'Smooth Hinge', 'Squared Hinge', 'Truncated SH', 'Truncated LS',
    'Smoothed Ramp1', 'Smoothed Ramp2', 'nonconvex exp loss1', 'nonconvex exp loss2', 'nonconvex exp loss3',
]

def save_loss_frequencies(loss_dictionary, dataset_name,file_name):

    if save == 0:
        return
    if noise == 1:
        scenario = "corrupted"
    elif non_iid == 1:
        scenario = "non_iid"
    else:
        scenario = "iid"
    
    sheet_name = f"{scenario}_loss_frequencies"
    loss_count = {loss_name: 0 for loss_name in loss_functions_ordered}

    for key, value in loss_dictionary.items():
        best_loss_data = value.get('best_loss')
        if best_loss_data:
            for loss_function in best_loss_data.keys():
                if loss_function in loss_count:
                    loss_count[loss_function] += 1

    message = f"Clients:{num_clients}, Parts:{participants}, CR:{CR}, Alpha:{alpha}"
    df_new = pd.DataFrame({
        'message': [message] * len(loss_functions_ordered),
        'datasetname': [dataset_name] * len(loss_functions_ordered),
        'Loss Function': loss_functions_ordered,
        'Frequency': [loss_count[loss] for loss in loss_functions_ordered]
    })

    filepath = file_name

    if not os.path.exists(filepath):
        with pd.ExcelWriter(filepath, engine='openpyxl') as writer:
            df_new.to_excel(writer, sheet_name=sheet_name, index=False)
    else:
        existing_book = load_workbook(filepath)
        with pd.ExcelWriter(filepath, engine='openpyxl', mode='a', if_sheet_exists='overlay') as writer:
            writer.workbook = existing_book
            if sheet_name in existing_book.sheetnames:
                start_row = existing_book[sheet_name].max_row
                df_new.to_excel(writer, sheet_name=sheet_name, index=False, header=False, startrow=start_row)
            else:
                df_new.to_excel(writer, sheet_name=sheet_name, index=False)


# ---------------------------------------
# Helper: map (noise, non_iid) ➜ suffix
# ---------------------------------------
def scenario_suffix(noise: int, non_iid: int) -> str:
    if noise == 0 and non_iid == 0:
        return "iid"
    if noise == 0 and non_iid == 1:
        return "noniid"
    if noise == 1 and non_iid == 0:
        return "noiseiid"
    return "noisenoniid"          # noise ==1 and non_iid ==1


# def heatmap_gen(multi_client_iter_data,single_client_iter_data,dataset):
#     # Prepare DataFrames
#     df_multi = pd.DataFrame(
#         multi_client_iter_data,
#         index=[f'Client_{i}' for i in range(num_clients)],
#         columns=[f'CR_{i+1}' for i in range(CR)]
#     )
#     df_single = pd.DataFrame(
#         single_client_iter_data,
#         index=[f'Client_{i}' for i in range(num_clients)],
#         columns=[f'CR_{i+1}' for i in range(CR)]
#     )

#     # Determine heatmap color scale max
#     max_value = max(df_multi.max().max(), df_single.max().max())

#     # Multi-loss heatmap
#     plt.figure(figsize=(12, 8))
#     ax = sns.heatmap(
#         df_multi, cmap='YlOrRd', annot=False,
#         vmin=0, vmax=max_value,
#         xticklabels=False, yticklabels=False,
#         cbar_kws={'label': 'Local Epochs'} # Keep the label here
#     )
#     cbar = ax.collections[0].colorbar

#     # Now, set the font size for the colorbar label
#     cbar.set_label('Local Epochs', fontsize=22) # <--- Correct way to set label font size

#     # If you also want to change the font size of the colorbar ticks:
#     cbar.ax.tick_params(labelsize=18)

#     plt.xlabel('Communication Rounds', fontsize=22)
#     plt.ylabel('Clients', fontsize=22)
#     # plt.title(f'{dataset} with Multi UniCSL')
#     plt.tight_layout()
#     if save:
#         plt.savefig(f'results/{(dataset)}_multi_loss_heatmap.png', dpi=200)
#     plt.show(block=True)

#     # Single-loss heatmap
#     plt.figure(figsize=(12, 8))
#     ax = sns.heatmap(
#         df_single, cmap='YlOrRd', annot=False,
#         vmin=0, vmax=max_value,
#         xticklabels=False, yticklabels=False,
#         cbar_kws={'label': 'Local Epochs'}
#     )
#     cbar = ax.collections[0].colorbar

#     # Now, set the font size for the colorbar label
#     cbar.set_label('Local Epochs', fontsize=20) # <--- Correct way to set label font size

#     # If you also want to change the font size of the colorbar ticks:
#     cbar.ax.tick_params(labelsize=14)

#     plt.xlabel('Communication Rounds', fontsize=20)
#     plt.ylabel('Clients',  fontsize=20)
#     # plt.title(f'{dataset} with Single UniCSL')
#     plt.tight_layout()
#     if save:
#         plt.savefig(f'results/{(dataset)}_single_loss_heatmap.png', dpi=200)
#     plt.show(block=True)

def plot_heatmap(df, dataset, suffix, max_value, font_sizes, xtick_interval, ytick_interval):
    plt.figure(figsize=(12, 8))
    ax = sns.heatmap(
        df, cmap='YlOrRd', annot=False,
        vmin=0, vmax=max_value,
        xticklabels=xtick_interval, yticklabels=ytick_interval,
        cbar_kws={'label': 'Local Epochs'}
    )
    cbar = ax.collections[0].colorbar
    ax.tick_params(axis='x', labelsize=18)
    ax.tick_params(axis='y', labelsize=18)
    cbar.set_label('Local Epochs', fontsize=font_sizes['cbar_label'])
    cbar.ax.tick_params(labelsize=font_sizes['cbar_ticks'])
    
    plt.xlabel('Communication Rounds', fontsize=font_sizes['xlabel'])
    plt.ylabel('Clients', fontsize=font_sizes['ylabel'])
    plt.tight_layout()
    if save:
        plt.savefig(f'results/{dataset}_{suffix}_heatmap.png', dpi=300)
    plt.close()

def heatmap_gen(multi_client_iter_data, single_client_iter_data, dataset):
    # Default font sizes
    font_sizes = {
        'cbar_label': 24,
        'cbar_ticks': 24,
        'xlabel': 24,
        'ylabel': 24
    }

    xtick_interval = 5
    ytick_interval = 10
    # Prepare DataFrames
    df_multi = pd.DataFrame(
        multi_client_iter_data,
        index=[f'Client_{i}' for i in range(num_clients)],
        columns=[f'{i}' for i in range(CR)]
    )
    df_single = pd.DataFrame(
        single_client_iter_data,
        index=[f'Client_{i}' for i in range(num_clients)],
        columns=[f'{i}' for i in range(CR)]
    )

    # Determine heatmap color scale max
    max_value = max(df_multi.max().max(), df_single.max().max())

    # Plot heatmaps
    plot_heatmap(df_multi, dataset, 'multi_loss', max_value, font_sizes, xtick_interval, ytick_interval)
    plot_heatmap(df_single, dataset, 'single_loss', max_value, font_sizes, xtick_interval, ytick_interval)



# old functions for old code structure
"""This code will store only global accuracies in an Excel file"""
def save_results_to_excel(filename, dataset_name, global_accuracy):
    message = f"client {num_clients}, participation {participants}, noniid {non_iid} local {local_epoch} cr {CR}"
    if save:
        # Prepare the data to be written to Excel
        data = {
            'dataset_name': [dataset_name] * len(global_accuracy),
            'message': [message] * len(global_accuracy),
            'round': list(range(1, len(global_accuracy) + 1)),
            'Global acc' : global_accuracy
        }
        
        # Convert the data into a DataFrame
        df = pd.DataFrame(data)
        
        # Check if the file already exists
        if os.path.exists(filename):
            # If it exists, append the data to the existing Excel file
            with pd.ExcelWriter(filename, mode='a', engine='openpyxl', if_sheet_exists='overlay') as writer:
                df.to_excel(writer, index=False, header=False, startrow=writer.sheets['Sheet1'].max_row)
        else:
            # If it does not exist, create a new Excel file and write the data
            with pd.ExcelWriter(filename, mode='w', engine='openpyxl') as writer:
                df.to_excel(writer, index=False)
        
        print(f"Results saved to {filename}")
    else:
        pass


def save_results_V1_V2(
        single_acc, multi_acc,
        single_time, single_mem,
        multi_time, multi_mem,
        single_iter_per_cr, multi_iter_per_cr,
        dataset,
        filename="results_v1_vs_v2.xlsx"
):
    
    if not save:
        return

    import os
    import pandas as pd

    message = f"clients {num_clients}, part {participants}, noniid {non_iid}, local {local_epoch}, cr {CR}"

    total_iter_v1 = sum(single_iter_per_cr)
    total_iter_v2 = sum(multi_iter_per_cr)

    # 1️⃣ Accuracy
    df_acc = pd.DataFrame({
        "dataset": dataset,
        "message": message,
        "round": range(1, len(single_acc) + 1),
        "UniCSL_V1_acc": single_acc,
        "UniCSL_V2_acc": multi_acc
    })

    # 2️⃣ Time
    df_time = pd.DataFrame([{
        "dataset": dataset,
        "message": message,
        "V1_time_sec": single_time,
        "V2_time_sec": multi_time
    }])

    # 3️⃣ Memory
    df_mem = pd.DataFrame([{
        "dataset": dataset,
        "message": message,
        "V1_memory_MB": single_mem,
        "V2_memory_MB": multi_mem
    }])

    # 4️⃣ Total Iterations
    df_iter = pd.DataFrame([{
        "dataset": dataset,
        "message": message,
        "V1_total_iter": total_iter_v1,
        "V2_total_iter": total_iter_v2
    }])

    # ---------- WRITE / APPEND ----------
    if os.path.exists(filename):
        with pd.ExcelWriter(filename, mode='a', engine='openpyxl', if_sheet_exists='overlay') as writer:

            def append_or_create(sheet_name, df):
                if sheet_name in writer.sheets:
                    start_row = writer.sheets[sheet_name].max_row
                    df.to_excel(writer, sheet_name=sheet_name,
                                index=False, header=False, startrow=start_row)
                else:
                    df.to_excel(writer, sheet_name=sheet_name, index=False)

            append_or_create("Accuracy", df_acc)
            append_or_create("Time", df_time)
            append_or_create("Memory", df_mem)
            append_or_create("Total_Iterations", df_iter)

    else:
        with pd.ExcelWriter(filename, mode='w', engine='openpyxl') as writer:
            df_acc.to_excel(writer, sheet_name="Accuracy", index=False)
            df_time.to_excel(writer, sheet_name="Time", index=False)
            df_mem.to_excel(writer, sheet_name="Memory", index=False)
            df_iter.to_excel(writer, sheet_name="Total_Iterations", index=False)

    print(f"Results appended to {filename}")


def save_results_unicsl(results_list, loss_functions, dataset,
                     filename="results/results_unicsl.xlsx"):

    if not save:
        return

    import os
    import pandas as pd

    message = f"clients {num_clients}, part {participants}, noniid {non_iid}, local {local_epoch}, cr {CR}"

    # Names of models
    model_names = [list(loss.keys())[0] for loss in loss_functions]
    model_names.append("UniCSL")  # dynamic one

    # ---------------- ACCURACY ----------------
    acc_dict = {
        "dataset": dataset,
        "message": message,
        "round": range(1, len(results_list[0][0]) + 1)
    }

    for name, res in zip(model_names, results_list):
        acc_dict[name] = res[0]  # global_accuracy

    df_acc = pd.DataFrame(acc_dict)
    # ---- Add Average Accuracy Row ----
    avg_row = {
        "dataset": dataset,
        "message": message,
        "round": "Average"
    }

    for name in model_names:
        avg_row[name] = df_acc[name].mean()

    df_acc = pd.concat([df_acc, pd.DataFrame([avg_row])], ignore_index=True)

    # ---------------- TOTAL ITERATION ----------------
    df_iter = pd.DataFrame([{
        "dataset": dataset,
        "message": message,
        **{name: res[1] for name, res in zip(model_names, results_list)}
    }])

    # ---------------- TIME ----------------
    df_time = pd.DataFrame([{
        "dataset": dataset,
        "message": message,
        **{name: res[4] for name, res in zip(model_names, results_list)}
    }])

    # ---------------- MEMORY ----------------
    df_mem = pd.DataFrame([{
        "dataset": dataset,
        "message": message,
        **{name: res[5] for name, res in zip(model_names, results_list)}
    }])

    # ---------------- WRITE / APPEND ----------------
    def append_or_create(writer, sheet_name, df):
        if sheet_name in writer.sheets:
            start_row = writer.sheets[sheet_name].max_row
            df.to_excel(writer, sheet_name=sheet_name,
                        index=False, header=False, startrow=start_row)
        else:
            df.to_excel(writer, sheet_name=sheet_name, index=False)

    if os.path.exists(filename):
        with pd.ExcelWriter(filename, mode='a', engine='openpyxl',
                            if_sheet_exists='overlay') as writer:
            append_or_create(writer, "Accuracy", df_acc)
            append_or_create(writer, "Total_Iteration", df_iter)
            append_or_create(writer, "Time", df_time)
            append_or_create(writer, "Memory", df_mem)
    else:
        with pd.ExcelWriter(filename, mode='w', engine='openpyxl') as writer:
            df_acc.to_excel(writer, sheet_name="Accuracy", index=False)
            df_iter.to_excel(writer, sheet_name="Total_Iteration", index=False)
            df_time.to_excel(writer, sheet_name="Time", index=False)
            df_mem.to_excel(writer, sheet_name="Memory", index=False)

    print(f"Results appended to {filename}")



def heatmap_gen_unicsl(results_list, loss_functions, dataset):

    model_names = [list(loss.keys())[0] for loss in loss_functions]
    model_names.append("UniCSL")

    font_sizes = {
        'cbar_label': 24,
        'cbar_ticks': 24,
        'xlabel': 24,
        'ylabel': 24
    }

    xtick_interval = 5
    ytick_interval = 10

    # collect all client_iter_data
    all_data = [res[2] for res in results_list]

    # global max for consistent scale
    max_value = max([data.max() for data in all_data])

    for name, client_iter_data in zip(model_names, all_data):

        df = pd.DataFrame(
            client_iter_data,
            index=[f'Client_{i}' for i in range(num_clients)],
            columns=[f'{i}' for i in range(CR)]
        )

        plot_heatmap(
            df,
            dataset,
            name.replace(" ", "_"),
            max_value,
            font_sizes,
            xtick_interval,
            ytick_interval
        )