import pandas as pd
import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
import copy
import time

from config import *
from Base.utils.utils import setup_logger, set_seed, scenario_suffix, save_loss_frequencies
from Base.utils.load_data import load_datasets

from Base.baselines.fedavg import main_fedavg
from Base.baselines.fedprox import main_fedprox
from Base.baselines.scaffold import main_scaffold
from Base.baselines.feddyn import main_feddyn
from Base.baselines.fedopt import main_fedopt
from Base.UniCSL.unicsl_static import main_unicsl as main_unicsl_static
from Base.UniCSL.unicsl_dynamic import main_unicsl as main_unicsl_dynamic

# ── Config ────────────────────────────────────────────────────────────────────
LR_LIST = [0.9, 0.5, 0.1, 0.01]

ALGO_RUNNERS = {
    'fedavg':        lambda gp, cs, tr, te, lr: main_fedavg(gp, cs, tr, te[0], lr),
    'fedprox':       lambda gp, cs, tr, te, lr: main_fedprox(gp, cs, tr, te[0], lr),
    'scaffold':      lambda gp, cs, tr, te, lr: main_scaffold(gp, cs, tr, te[0], lr),
    'feddyn':        lambda gp, cs, tr, te, lr: main_feddyn(gp, cs, tr, te[0], lr),
    'fedopt':        lambda gp, cs, tr, te, lr: main_fedopt(gp, cs, tr, te[0], lr),
    'unicsl_static': lambda gp, cs, tr, te, lr: main_unicsl_static(gp, cs, tr, te, lr)[0],
    # 'unicsl_dynamic':lambda gp, cs, tr, te, lr: main_unicsl_dynamic(gp, cs, tr, te, lr),
}

set_seed()
log_file, tee = setup_logger(save_log=save_log)

client_schedule = [np.random.choice(num_clients, size=participants, replace=False) for _ in range(CR)]

# ── Results storage: {algo: {dataset: {lr: (avg_acc, std)}}} ─────────────────
results = {algo: {ds: {} for ds in all_datasets} for algo in aggregations}

# ── Main tuning loop ──────────────────────────────────────────────────────────
for dataset in all_datasets:
    train_data_list, global_test_data = load_datasets(dataset)
    global_parameters = np.random.rand(train_data_list[0].shape[1])

    for aggregation in aggregations:
        if aggregation not in ALGO_RUNNERS:
            print(f"[WARN] Unknown aggregation '{aggregation}', skipping.")
            continue

        runner = ALGO_RUNNERS[aggregation]

        for lr in LR_LIST:
            print(f"\n[RUN] dataset={dataset} | algo={aggregation} | lr={lr}")
            try:
                result = runner(
                    copy.deepcopy(global_parameters),
                    client_schedule,
                    copy.deepcopy(train_data_list),
                    copy.deepcopy(global_test_data),
                    lr
                )
                acc_array = np.array(result[0])
                avg_acc = float(np.mean(acc_array))
                std_acc = float(np.std(acc_array))
                print(f"  → avg_acc={avg_acc:.4f} ± {std_acc:.4f}  | time={result[1]:.1f}s")
            except Exception as e:
                print(f"  [ERROR] {e}")
                avg_acc, std_acc = float('nan'), float('nan')

            results[aggregation][dataset][lr] = (avg_acc, std_acc)

# ── Write to Excel ─────────────────────────────────────────────────────────────
def make_excel(results, lr_list, output_path=f"lr_rate_noniid_{non_iid}_noise_{noise}.xlsx"):
    wb = Workbook()

    header_fill  = PatternFill("solid", start_color="2F75B6", end_color="2F75B6")
    header_font  = Font(bold=True, color="FFFFFF", name="Arial")
    cell_font    = Font(name="Arial", size=10)
    center       = Alignment(horizontal="center", vertical="center")
    thin         = Side(style="thin", color="AAAAAA")
    border       = Border(left=thin, right=thin, top=thin, bottom=thin)

    first = True
    for algo in results:
        ws = wb.active if first else wb.create_sheet()
        ws.title = algo
        first = False

        # Header row
        ws.cell(1, 1, "dataset").font = header_font
        ws.cell(1, 1).fill = header_fill
        ws.cell(1, 1).alignment = center
        ws.cell(1, 1).border = border
        ws.column_dimensions['A'].width = 18

        for col_idx, lr in enumerate(lr_list, start=2):
            cell = ws.cell(1, col_idx, f"lr {lr}")
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = center
            cell.border = border
            from openpyxl.utils import get_column_letter
            ws.column_dimensions[get_column_letter(col_idx)].width = 20

        # Data rows
        for row_idx, dataset in enumerate(results[algo], start=2):
            ws.cell(row_idx, 1, dataset).font = Font(bold=True, name="Arial")
            ws.cell(row_idx, 1).alignment = center
            ws.cell(row_idx, 1).border = border

            for col_idx, lr in enumerate(lr_list, start=2):
                avg, std = results[algo][dataset].get(lr, (float('nan'), float('nan')))
                if not (np.isnan(avg) or np.isnan(std)):
                    value = f"{avg:.4f} ± {std:.4f}"
                else:
                    value = "ERROR"
                cell = ws.cell(row_idx, col_idx, value)
                cell.font = cell_font
                cell.alignment = center
                cell.border = border

        ws.row_dimensions[1].height = 20

    wb.save(output_path)
    print(f"\n[DONE] Results saved to '{output_path}'")

make_excel(results, LR_LIST, output_path=f"results/lr_rate_noniid_{non_iid}_noise_{noise}.xlsx")

if log_file:
    log_file.close()
    tee.close()